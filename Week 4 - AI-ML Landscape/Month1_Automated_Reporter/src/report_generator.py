from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime
from openpyxl.drawing.image import Image

def create_report(customer_df, reviews_df, customer_summary, reviews_summary):

    """
    Creates the Excel report.
    """

    # Create workbook
    wb = Workbook()

    # ==========================
    # Executive Summary Sheet
    # ==========================
    ws = wb.active
    ws.title = "Executive Summary"

    # Report Title
    ws["A1"] = "CUSTOMER INSIGHTS PERFORMANCE REPORT"
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(
        start_color="7030A0",
        end_color="7030A0",
        fill_type="solid"
    )
    ws.merge_cells("A1:B1")
    ws["A1"].alignment = Alignment(horizontal="center")

    # Header Style
    header_fill = PatternFill(
        start_color="4472C4",
        end_color="4472C4",
        fill_type="solid"
    )

    ws["A3"] = "Customer Insight"
    ws["B3"] = "Value"

    for cell in ["A3", "B3"]:
        ws[cell].fill = header_fill
        ws[cell].font = Font(bold=True, color="FFFFFF")

    # Customer Summary
    ws.append([
        "Report Generated",
        datetime.now().strftime("%d %B %Y %I:%M %p")
    ])

    ws.append([
        "Total Customers",
        customer_summary["total_customers"]
    ])

    ws.append([
        "Average Response Time (Hours)",
        customer_summary["average_response_time"]
    ])

    ws.append([
        "Top Region",
        customer_summary["top_region"]
    ])

    ws.append([
        "Issue Resolution Rate",
        customer_summary["resolution_rate"] / 100
    ])

    # Review Summary
    ws.append([
        "Total Reviews",
        reviews_summary["total_reviews"]
    ])

    ws.append([
        "Average Rating",
        reviews_summary["average_rating"]
    ])

    ws.append([
        "Top Review Location",
        reviews_summary["top_location"]
    ])

    # Format percentage
    ws["B8"].number_format = "0.00%"

    # Column widths
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 28

           # ==========================
    # Customer Data Sheet
    # ==========================

    ws2 = wb.create_sheet("Customer Data")
    ws2.freeze_panes = "A2"

    ws2.append(customer_df.columns.tolist())

    for cell in ws2[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")

    for row in customer_df.values.tolist():
        ws2.append(row)

    # Auto-adjust column widths
    for column in ws2.columns:
        length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column
        )
        ws2.column_dimensions[column[0].column_letter].width = min(length + 3, 35) 

            # ==========================
    # Reviews Data Sheet
    # ==========================

    ws3 = wb.create_sheet("Reviews Data")
    ws3.freeze_panes = "A2"

    ws3.append(reviews_df.columns.tolist())

    for cell in ws3[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")

    for row in reviews_df.values.tolist():
        ws3.append(row)

    # Auto-adjust column widths
    for column in ws3.columns:
        length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column
        )
        ws3.column_dimensions[column[0].column_letter].width = min(length + 3, 35)

            # ==========================
    # Charts Sheet
    # ==========================

    charts_sheet = wb.create_sheet("Charts")

    chart_files = [
        "output/sentiment_distribution.png",
        "output/rating_distribution.png",
        "output/customers_by_region.png",
        "output/issue_resolution.png",
        "output/response_time_distribution.png"
    ]

    positions = ["A1", "J1", "A20", "J20", "A39"]

    for file, position in zip(chart_files, positions):
        try:
            img = Image(file)
            img.width = 450
            img.height = 300
            charts_sheet.add_image(img, position)
        except Exception as e:
            print(f"Could not add {file}: {e}")

    # ==========================
    # Save Report
    # ==========================

    wb.save("output/Customer_Insights_Report.xlsx")

    print("✅ Customer Insights Excel report created successfully!")