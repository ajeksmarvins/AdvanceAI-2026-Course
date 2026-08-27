import pandas as pd
import matplotlib.pyplot as plt
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side
)

from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.drawing.image import Image

print("=" * 50)
print("WEEK 2 REPORT GENERATOR")
print("=" * 50)

try:
    df = pd.read_csv("sales_data_CLEANED.csv")
    print("✅ Data loaded successfully.")
except FileNotFoundError:
    print("❌ sales_data_CLEANED.csv not found.")
    exit()

# Create total_sales if it doesn't exist
if "total_sales" not in df.columns:
    df["total_sales"] = df["quantity"] * df["unit_price"]

print(f"Rows Loaded : {len(df)}")
print(f"Columns : {len(df.columns)}")
# -----------------------------
# CREATE CHARTS
# -----------------------------

print("\nCreating charts...")

# Sales by Region
region_sales = df.groupby("region")["total_sales"].sum()

plt.figure(figsize=(7,5))
region_sales.plot(kind="bar", color="steelblue")
plt.title("Total Sales by Region")
plt.xlabel("Region")
plt.ylabel("Sales")
plt.tight_layout()
plt.savefig("chart_region.png")
plt.close()


# Sales by Product
product_sales = (
    df.groupby("product")["total_sales"]
      .sum()
      .sort_values(ascending=False)
)

plt.figure(figsize=(8,5))
product_sales.plot(kind="bar", color="green")
plt.title("Sales by Product")
plt.xlabel("Product")
plt.ylabel("Sales")
plt.tight_layout()
plt.savefig("chart_product.png")
plt.close()


# Quantity Distribution
qty = df.groupby("product")["quantity"].sum()

plt.figure(figsize=(6,6))
qty.plot(kind="pie", autopct="%1.1f%%")
plt.ylabel("")
plt.title("Product Quantity Distribution")
plt.tight_layout()
plt.savefig("chart_quantity.png")
plt.close()


# Unit Price Distribution
plt.figure(figsize=(7,5))
plt.hist(df["unit_price"], bins=10)
plt.title("Unit Price Distribution")
plt.xlabel("Price")
plt.ylabel("Frequency")
plt.tight_layout()
plt.savefig("chart_price.png")
plt.close()

print("✅ Charts created successfully.")
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.drawing.image import Image
from openpyxl.utils import get_column_letter

print("\nCreating Excel Report...")

wb = Workbook()

# =========================
# Executive Summary Sheet
# =========================
ws1 = wb.active
ws1.title = "Executive Summary"

title_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")

title_font = Font(size=16, bold=True, color="FFFFFF")
header_font = Font(bold=True, color="FFFFFF")

# Title
ws1.merge_cells("A1:B1")
ws1["A1"] = "WEEKLY SALES PERFORMANCE REPORT"
ws1["A1"].font = title_font
ws1["A1"].fill = title_fill
ws1["A1"].alignment = Alignment(horizontal="center")

# Headers
ws1["A3"] = "Metric"
ws1["B3"] = "Value"

for cell in ["A3", "B3"]:
    ws1[cell].font = header_font
    ws1[cell].fill = header_fill

# Metrics
ws1.append(["Report Generated", datetime.now().strftime("%d %B %Y %I:%M %p")])
ws1.append(["Total Sales", df["total_sales"].sum()])
ws1.append(["Total Orders", len(df)])
ws1.append(["Average Sale", round(df["total_sales"].mean(), 2)])
# Format currency cells
ws1["B5"].number_format = "$#,##0.00"
ws1["B7"].number_format = "$#,##0.00"

# Column Width
ws1.column_dimensions["A"].width = 25
ws1.column_dimensions["B"].width = 18

# Insert Charts
try:
    img1 = Image("chart_region.png")
    img1.width = 350
    img1.height = 250
    ws1.add_image(img1, "D2")

    img2 = Image("chart_product.png")
    img2.width = 350
    img2.height = 250
    ws1.add_image(img2, "D18")
except:
    print("Charts not inserted.")

# =========================
# Detailed Breakdown Sheet
# =========================
ws2 = wb.create_sheet("Detailed Breakdown")

for row in dataframe_to_rows(df, index=False, header=True):
    ws2.append(row)

for cell in ws2[1]:
    cell.font = header_font
    cell.fill = header_fill

# Auto-size columns
for column_cells in ws2.columns:
    length = max(len(str(cell.value)) if cell.value else 0 for cell in column_cells)
    ws2.column_dimensions[get_column_letter(column_cells[0].column)].width = length + 3

# Insert remaining charts
try:
    img3 = Image("chart_quantity.png")
    img3.width = 350
    img3.height = 250
    ws2.add_image(img3, "J2")

    img4 = Image("chart_price.png")
    img4.width = 350
    img4.height = 250
    ws2.add_image(img4, "J20")
except:
    print("Additional charts not inserted.")

# Save workbook
wb.save("Sales_Report.xlsx")

print("✅ Excel report created successfully!")